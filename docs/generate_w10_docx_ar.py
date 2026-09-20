import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_paragraph_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    pPr = p._p.get_or_add_pPr()
    bidi = parse_xml(f'<w:bidi {nsdecls("w")}/>')
    pPr.append(bidi)
    p.alignment = align

def set_run_rtl(run, font_name="Segoe UI"):
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
    rtl = parse_xml(f'<w:rtl {nsdecls("w")}/>')
    rPr.append(rFonts)
    rPr.append(rtl)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    # Set table level RTL
    bidiVisual = parse_xml(f'<w:bidiVisual {nsdecls("w")}/>')
    tblPr.append(bidiVisual)

def add_rtl_paragraph(doc, text="", space_after=4, bold=False, size=10.5, color=RGBColor(0x1F, 0x29, 0x37), align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = doc.add_paragraph()
    set_paragraph_rtl(p, align)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.2
    if text:
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        set_run_rtl(r, "Segoe UI")
    return p

def add_rtl_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Segoe UI"
    r.font.bold = True
    set_run_rtl(r, "Segoe UI")
    
    if level == 1:
        r.font.size = Pt(15)
        r.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B) # Deep indigo
    elif level == 2:
        r.font.size = Pt(12.5)
        r.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5) # Brand indigo
    else:
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    return p

def build_arabic_document():
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # -------------------------------------------------------------
    # 1. صفحة الغلاف الرسمية (Cover Page)
    # -------------------------------------------------------------
    p_inst = doc.add_paragraph()
    set_paragraph_rtl(p_inst, WD_ALIGN_PARAGRAPH.CENTER)
    p_inst.paragraph_format.space_after = Pt(20)
    
    r_inst = p_inst.add_run(
        "الجمهورية اليمنية\n"
        "وزارة التعليم العالي والبحث العلمي\n"
        "جامعة الحكمة - صنعاء\n"
        "كلية الهندسة وتكنولوجيا المعلومات\n"
        "قسم تكنولوجيا المعلومات (IT)\n"
    )
    r_inst.font.name = "Segoe UI"
    r_inst.font.size = Pt(11.5)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    set_run_rtl(r_inst)

    p_title = doc.add_paragraph()
    set_paragraph_rtl(p_title, WD_ALIGN_PARAGRAPH.CENTER)
    p_title.paragraph_format.space_before = Pt(20)
    p_title.paragraph_format.space_after = Pt(24)

    r_t1 = p_title.add_run("التقرير النهائي الشامل لتوثيق المشروع البرمجي\n")
    r_t1.font.name = "Segoe UI"
    r_t1.font.size = Pt(20)
    r_t1.font.bold = True
    r_t1.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B)
    set_run_rtl(r_t1)

    r_t2 = p_title.add_run("نظام حجز وإدارة الفنادق المتكامل (Enterprise Hotel Booking System)\n")
    r_t2.font.name = "Segoe UI"
    r_t2.font.size = Pt(14)
    r_t2.font.bold = True
    r_t2.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
    set_run_rtl(r_t2)

    r_t3 = p_title.add_run("دليل التدريب الميداني — محاكاة هندسة البرمجيات الشاملة (2026 - 2027)\n")
    r_t3.font.name = "Segoe UI"
    r_t3.font.size = Pt(11)
    r_t3.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    set_run_rtl(r_t3)

    # جدول بيانات الغلاف
    meta_table = doc.add_table(rows=8, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table, "E5E7EB", "6")

    meta_items = [
        ("عنوان المشروع البرمجي", "نظام حجز وإدارة الفنادق متعدد الطبقات (Hotel Booking System)"),
        ("المساق الأكاديمي", "دليل التدريب الميداني لمشاريع Full-Stack الفردية (IT-492)"),
        ("المشرف الأكاديمي", "أ.د / إبراهيم أحمد البلطة (أستاذ هندسة البرمجيات وتكنولوجيا المعلومات)"),
        ("إعداد مهندس البرمجيات", "لؤي السروري (Loai Alsorory)"),
        ("النمط المعماري للمشروع", "معمارية البصلة / النظيفة (Clean Architecture) + Repository Pattern"),
        ("مستودع الباك إند والويب (GitHub)", "https://github.com/LoaiAlsorory/HotelBookingSystem"),
        ("مستودع تطبيق الموبايل (Flutter)", "https://github.com/LoaiAlsorory/HotelBookingSystem-Flutter"),
        ("الإصدار وتاريخ التقديم", "الإصدار النهائي 1.0.0 — سبتمبر 2026م")
    ]

    for idx, (label, val) in enumerate(meta_items):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.5)
        cell_val.width = Inches(4.0)

        set_cell_background(cell_lbl, "F3F4F6")
        set_cell_margins(cell_lbl, 80, 80, 120, 120)
        set_cell_margins(cell_val, 80, 80, 120, 120)

        p0 = cell_lbl.paragraphs[0]
        set_paragraph_rtl(p0, WD_ALIGN_PARAGRAPH.RIGHT)
        r0 = p0.add_run(label)
        r0.font.name = "Segoe UI"
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        set_run_rtl(r0)

        p1 = cell_val.paragraphs[0]
        set_paragraph_rtl(p1, WD_ALIGN_PARAGRAPH.RIGHT)
        r1 = p1.add_run(val)
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        set_run_rtl(r1)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 2. الملخص التنفيذي والمعمارية العامة (Executive Summary)
    # -------------------------------------------------------------
    add_rtl_heading(doc, "1. الملخص التنفيذي والمعمارية البرمجية للمشروع", level=1)
    
    add_rtl_paragraph(
        doc,
        "يُعد هذا التقرير التوثيقي المرجع الفني والأكاديمي النهائي لتسليم متطلبات مساق التدريب الميداني بجامعة الحكمة، "
        "المُقام تحت إشراف أ.د/ إبراهيم أحمد البلطة. تم تصميم وتنفيذ المشروع بمحاكاة دقيقة لبيئات شركات البرمجيات العالمية، "
        "حيث تولى الطالب دور مهندس برمجيات متكامل (Full-Stack Software Engineer) لتنفيذ كافة مراحل دورة حياة تطوير النظم (SDLC) "
        "بشكل فردي وصارم."
    )
    
    add_rtl_paragraph(
        doc,
        "يتكون النظام من منصة سحابية شاملة لإدارة وحجز الفنادق، مبنية وفق أحدث الممارسات الهندسية وتتوزع على مشروعين منفصلين تماماً:"
    )

    p = add_rtl_paragraph(doc, "", space_after=2)
    r1 = p.add_run("أولاً: مشروع الباك إند وبوابة الويب الإدارية (HotelBookingSystem): ")
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
    set_run_rtl(r1)
    r2 = p.add_run("مبني بإطار العمل ASP.NET Core 10 وفق نمط Clean Architecture ونمط المستودع Repository Pattern، ويحتوي قاعدة البيانات العلائقية ومحرك واجهات البرمجة RESTful APIs وبوابة الإدارة بتقنية ASP.NET Core MVC.")
    set_run_rtl(r2)

    p = add_rtl_paragraph(doc, "", space_after=6)
    r1 = p.add_run("ثانياً: تطبيق الهاتف المحمول المستقل (HotelBookingSystem-Flutter): ")
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x7C, 0x3A, 0xED)
    set_run_rtl(r1)
    r2 = p.add_run("مبني بإطار العمل Flutter ومفصول في مشروع ومستودع مستقل كلياً تنفيذاً لتوجيهات المشرف الأكاديمي، ويعمل كعميل مستقل (Decoupled Client) يستهلك خدمات الـ Web API عبر بروتوكول HTTP الآمن.")
    set_run_rtl(r2)

    # هيكل الطبقات البرمجية
    add_rtl_heading(doc, "مخطط سريان الاعتمادية المعمارية (Clean Architecture):", level=3)
    p_code = doc.add_paragraph()
    p_code.paragraph_format.space_before = Pt(4)
    p_code.paragraph_format.space_after = Pt(8)
    r = p_code.add_run(
        "┌────────────────────────────────┐         ┌────────────────────────────────┐\n"
        "│  بوابة الويب الإدارية (MVC)    │         │  تطبيق الهاتف المحمول (Flutter)│\n"
        "│  (Bootstrap 5 RTL + DataTable) │         │  (Provider Pattern + Cairo)    │\n"
        "└───────────────┬────────────────┘         └───────────────┬────────────────┘\n"
        "                │                                          │ استدعاء برمجيات (REST/JSON)\n"
        "                ▼                                          ▼\n"
        "     [ HotelBookingSystem.Application ] ◄─────── [ HotelBookingSystem.API ]\n"
        "                │\n"
        "                ▼\n"
        "     [ HotelBookingSystem.Domain ] (الكيانات وقواعد العمل الخالصة)\n"
        "                ▲\n"
        "                │ تزويد البيانات عبر Entity Framework Core 10\n"
        "     [ HotelBookingSystem.Infrastructure ]\n"
        "                │\n"
        "                ▼\n"
        "     [ قاعدة بيانات SQL Server ]"
    )
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # -------------------------------------------------------------
    # 3. الجزء الأول: دليل التثبيت والتجهيز البيئي (W10.1)
    # -------------------------------------------------------------
    add_rtl_heading(doc, "2. الجزء الأول: دليل التثبيت والتجهيز البيئي والتشغيل (W10.1 Installation Guide)", level=1)

    add_rtl_paragraph(
        doc,
        "يحدد هذا القسم المتطلبات الهندسية والخطوات الإجرائية الدقيقة لتجهيز بيئة العمل وتشغيل كافة مكونات النظام من الصفر."
    )

    add_rtl_heading(doc, "1.2 المتطلبات المسبقة لتشغيل النظام (Prerequisites):", level=2)

    table_req = doc.add_table(rows=6, cols=3)
    table_req.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_req, "E5E7EB", "4")

    headers = ["المكون البرمجي / التقنية", "الحد الأدنى المطلوب", "الغرض في النظام"]
    for i, h in enumerate(headers):
        cell = table_req.rows[0].cells[i]
        set_cell_background(cell, "1E1B4B")
        p = cell.paragraphs[0]
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run(h)
        r.font.name = "Segoe UI"
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)
        set_run_rtl(r)

    req_data = [
        (".NET SDK", "الإصدار 10.0 (x64)", "ترجمة وتشغيل مشروعي Web API وبوابة الويب MVC"),
        ("محرك قاعدة البيانات", "SQL Server 2022 أو LocalDB", "حفظ واسترجاع بيانات الفنادق، الغرف، الحجوزات، والمستخدمين"),
        ("إطار عمل الموبايل", "Flutter SDK 3.19 فما فوق مع Dart 3", "بناء وتشغيل تطبيق الهاتف المحمول عبر المتصفح أو أجهزة أندرويد"),
        ("بيئة التطوير والمحررات", "Visual Studio 2022 / VS Code", "تحرير الأكواد وإدارة الحلول البرمجية والمشروعات"),
        ("نظام إدارة الإصدارات", "Git CLI 2.40+", "مزامنة التحديثات، إدارة الفروع، والرفع إلى مستودعات GitHub")
    ]

    for row_idx, data in enumerate(req_data, start=1):
        row = table_req.rows[row_idx]
        bg = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 100, 100)
            p = cell.paragraphs[0]
            set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            r = p.add_run(text)
            r.font.name = "Segoe UI"
            r.font.size = Pt(9)
            set_run_rtl(r)

    add_rtl_heading(doc, "2.2 إعداد قاعدة البيانات والزرع التلقائي للبيانات (Automated Seeding):", level=2)
    add_rtl_paragraph(
        doc,
        "تم تضمين منطق تهيئة ذاتي لقاعدة البيانات داخل ملف Program.cs؛ حيث يقوم النظام تلقائياً عند أول تشغيل بما يلي:\n"
        "1. استدعاء Database.EnsureCreated() للتحقق من وجود قاعدة البيانات وبناء الجداول والعلاقات دون الحاجة لكتابة أوامر SQL يدوية.\n"
        "2. تشغيل كلاس الزرع DbSeeder.SeedAsync() لملء قاعدة البيانات فوراً ببيانات تجريبية واقعية تشمل:\n"
        "   - 12 فندقاً فاخراً فريداً في عواصم ومدن عربية كبرى، مع صور حقيقية عالية الجودة لكل فندق على حدة.\n"
        "   - 48 غرفة فندقية موزعة على الفنادق، تختلف في السعات والأسعار وتملك كل منها صورة غرفة مستقلة.\n"
        "   - حسابات مستخدمين مشفرة بكلمات مرور محمية بخوارزمية PBKDF2 الموصى بها عالمياً.\n"
        "   - نماذج حجوزات مسبقة لاختبار دورة حياة الحجز وتأكيد التوفر."
    )

    add_rtl_heading(doc, "3.2 خطوات وأوامر تشغيل المشاريع (Execution Commands):", level=2)
    
    p_code = doc.add_paragraph()
    p_code.paragraph_format.space_before = Pt(4)
    p_code.paragraph_format.space_after = Pt(8)
    r = p_code.add_run(
        "# الخطوة 1: استنساخ وبناء مشروع الباك إند والويب:\n"
        "git clone https://github.com/LoaiAlsorory/HotelBookingSystem.git\n"
        "cd HotelBookingSystem\n"
        "dotnet restore HotelBookingSystem.sln\n"
        "dotnet build HotelBookingSystem.sln\n\n"
        "# الخطوة 2: تشغيل خدمة الـ Web API (في نافذة Terminal مستقلة):\n"
        "cd src/HotelBookingSystem.API\n"
        "dotnet run\n"
        "# الرابط المباشر: http://localhost:5242\n"
        "# رابط توثيق Swagger التفاعلي: http://localhost:5242/swagger\n\n"
        "# الخطوة 3: تشغيل بوابة الويب الإدارية MVC (في نافذة Terminal ثانية):\n"
        "cd src/HotelBookingSystem.Web\n"
        "dotnet run\n"
        "# الرابط المباشر للبوابة: http://localhost:5200\n"
        "# بيانات تسجيل الدخول الافتراضية للمدير:\n"
        "# البريد: admin@hotel.com  |  كلمة المرور: Password123!\n\n"
        "# الخطوة 4: تشغيل تطبيق الهاتف المحمول Flutter (في نافذة Terminal ثالثة):\n"
        "cd E:/Project_ASP.NET/HOTEL_BOOKIN_SYSTEM/Hotel_Booking_Flutter\n"
        "flutter pub get\n"
        "flutter run -d chrome     # أو flutter run للأجهزة الحقيقية المربوطة بالـ USB"
    )
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 4. الجزء الثاني: دليل المستخدم المصور (W10.2 User Manual)
    # -------------------------------------------------------------
    add_rtl_heading(doc, "3. الجزء الثاني: دليل المستخدم الشامل وواجهات النظام (W10.2 User Manual)", level=1)

    add_rtl_paragraph(
        doc,
        "يقدم النظام واجهتين متكاملتين: الأولى بوابة إدارة ويب مخصصة للمدراء وموظفي الحجز (Web Portal)، "
        "والثانية تطبيق هاتف محمول سلس مخصص للنزلاء والعملاء (Flutter Mobile App)."
    )

    add_rtl_heading(doc, "1.3 دليل استخدام بوابة الويب الإدارية (ASP.NET Core MVC):", level=2)

    web_features = [
        ("واجهة تسجيل الدخول المنقسمة (Split-Screen Login)", "تعتمد تصميماً هندسياً فاخراً يجمع بين عرض صور الفنادق على اليمين ونموذج تسجيل الدخول الذكي على اليسار، مع فحص صحة المدخلات وإصدار ملف تعريف ارتباط مشفر (Secure Cookie)."),
        ("لوحة التحكم المركزية (Main Dashboard)", "تعرض 4 بطاقات إحصائيات تفاعلية (إجمالي الفنادق، الغرف الشاغرة، الحجوزات النشطة، إجمالي الإيرادات المالية المحصلة)، بالإضافة لجدول أحدث 5 حجوزات مع شارات الحالة الملونة وأزرار الوصول السريع."),
        ("إدارة الفنادق (Hotels Management)", "تتيح استعراض الفنادق في جدول تفاعلي مدعوم بـ DataTable.js للبحث اللحظي والفرز. يمكن إضافة فندق، تعديل بياناته، أو حذفه بأمان عبر نافذة تأكيد منبثقة (Bootstrap Modal)."),
        ("إدارة الغرف الفندقية (Rooms Inventory)", "تمكين الموظف من ربط الغرف بالفندق، وتحديد نوع الغرفة (مفردة، مزدوجة، جناح، ديلوكس)، وتحديد طاقتها الاستيعابية، وسعر الليلة، وتعديل حالة التوفر بضغطة زر."),
        ("إدارة دورة حياة الحجوزات (Bookings Management)", "إمكانية إنشاء حجز فندقي جديد مع احتساب السعر الإجمالي تلقائياً (عدد الليالي × سعر الليلة)، والتحقق من عدم وجود حجز سابق في نفس الفترة، مع خيارات تأكيد أو إلغاء الحجز."),
        ("دليل المستخدمين والصلاحيات (Users Management)", "استعراض قائمة المستخدمين والعملاء المسجلين، وأدوارهم وصلاحياتهم في النظام.")
    ]

    for title, desc in web_features:
        p = add_rtl_paragraph(doc, "", space_after=3)
        r_t = p.add_run(f"• {title}: ")
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
        set_run_rtl(r_t)
        r_d = p.add_run(desc)
        set_run_rtl(r_d)

    add_rtl_heading(doc, "2.3 دليل استخدام تطبيق الهاتف المحمول (Flutter Mobile App):", level=2)

    mobile_features = [
        ("شاشة الترحيب والدخول (Splash & Auth)", "عرض شعار النظام مع التحقق التلقائي من حالة تسجيل الدخول المخزنة، ونموذج دخول ذكي مع إمكانية إخفاء وإظهار كلمة المرور."),
        ("الشاشة الرئيسية (HomeScreen Dashboard)", "بطاقات إحصائية للمستخدم، شريط ترحيبي باسم المستخدم، شريط أفقي متحرك لأبرز الفنادق، وأزرار الانتقال المباشر."),
        ("القائمة الجانبية (Drawer)", "قائمة سحب من اليمين متوافقة مع الـ RTL تحتوي على صورة المستخدم، بريده، صلاحيته، وروابط الانتقال وزر تسجيل الخروج."),
        ("شريط التنقل السفلي (Bottom Navigation Bar)", "شريط ثابت يضم 5 تبويبات أساسية (لوحة التحكم، المستخدمون، الفنادق، الغرف، الحجوزات) مع تمييز التبويب النشط بخط Cairo."),
        ("استعراض الفنادق وتفاصيلها (Hotels & Details)", "شاشة كتالوج الفنادق مع شريط بحث بالاسم والمدينة، وشاشة تفاصيل الفندق الفاخرة التي تستعرض المرافق وقائمة الغرف التابعة له."),
        ("النافذة المنبثقة لحجز الغرفة (SheetModal)", "عند الضغط على 'حجز الآن' على أي غرفة، تنبثق نافذة سفلية تفاعلية تمكن المستخدم من تحديد تواريخ الوصول والمغادرة وعدد الضيوف ومراجعة السعر الإجمالي وتأكيد الحجز مع إشعار نجاح فوري.")
    ]

    for title, desc in mobile_features:
        p = add_rtl_paragraph(doc, "", space_after=3)
        r_t = p.add_run(f"• {title}: ")
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(0x7C, 0x3A, 0xED)
        set_run_rtl(r_t)
        r_d = p.add_run(desc)
        set_run_rtl(r_d)

    # -------------------------------------------------------------
    # 5. الجزء الثالث: تقرير الفحص وضمان الجودة (W10.3 Testing)
    # -------------------------------------------------------------
    add_rtl_heading(doc, "4. الجزء الثالث: تقرير الفحص والاختبارات وضمان الجودة (W10.3 Testing & QA)", level=1)

    add_rtl_paragraph(
        doc,
        "خضع النظام لمصفوفة اختبارات برمجية شاملة شملت واجهات الـ RESTful APIs، التحقق من صحة المدخلات، "
        "وقواعد منطق الأعمال (Business Invariants)."
    )

    add_rtl_heading(doc, "1.4 نتائج اختبارات واجهات الـ Web API عبر Postman و Swagger:", level=2)

    table_tests = doc.add_table(rows=9, cols=5)
    table_tests.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_tests, "E5E7EB", "4")

    cols_test = ["نقطة النهاية (Endpoint)", "نوع الطلب", "سيناريو الاختبار", "رمز الاستجابة المتوقع", "النتيجة الفعلية"]
    for i, c in enumerate(cols_test):
        cell = table_tests.rows[0].cells[i]
        set_cell_background(cell, "1E1B4B")
        p = cell.paragraphs[0]
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run(c)
        r.font.name = "Segoe UI"
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9)
        set_run_rtl(r)

    tests_data = [
        ("/api/auth/login", "POST", "تسجيل دخول صحيح للمدير", "200 OK", "ناجح (استلام JWT Token)"),
        ("/api/auth/login", "POST", "محاولة دخول بكلمة مرور خاطئة", "401 Unauthorized", "ناجح (رفض وتنبيه بالخطأ)"),
        ("/api/hotels", "GET", "استرجاع قائمة الفنادق كاملة", "200 OK", "ناجح (استرجاع 12 فندقاً)"),
        ("/api/hotels/99999", "GET", "طلب فندق برقم غير موجود", "404 Not Found", "ناجح (إرجاع رسالة غير موجود)"),
        ("/api/hotels", "POST", "إضافة فندق بدون اسم (حقل إجباري)", "400 Bad Request", "ناجح (رفض مع رسالة التحقق)"),
        ("/api/rooms", "GET", "تصفية الغرف التابعة لفندق محدد", "200 OK", "ناجح (استرجاع غرف الفندق)"),
        ("/api/bookings", "POST", "إنشاء حجز فندقي بتاريخ صحيح", "201 Created", "ناجح (إنشاء الحجز وحفظه)"),
        ("/api/bookings", "POST", "محاولة حجز مزدوج لغرفة محجوزة", "409 Conflict", "ناجح (منع الحجز المزدوج)")
    ]

    for row_idx, data in enumerate(tests_data, start=1):
        row = table_tests.rows[row_idx]
        bg = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 50, 50, 80, 80)
            p = cell.paragraphs[0]
            set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            r = p.add_run(val)
            r.font.name = "Segoe UI"
            r.font.size = Pt(8.5)
            if col_idx == 4:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x05, 0x96, 0x69)
            set_run_rtl(r)

    add_rtl_heading(doc, "2.4 مصفوفة التحقق من صحة المدخلات (Input Validation Matrix):", level=2)

    table_val = doc.add_table(rows=6, cols=4)
    table_val.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_val, "E5E7EB", "4")

    cols_val = ["الكيان / الحقل", "قاعدة التحقق البرمجية", "حالة الفحص الخاطئة (Negative Test)", "استجابة النظام وتصرفه"]
    for i, c in enumerate(cols_val):
        cell = table_val.rows[0].cells[i]
        set_cell_background(cell, "374151")
        p = cell.paragraphs[0]
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run(c)
        r.font.name = "Segoe UI"
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9)
        set_run_rtl(r)

    val_data = [
        ("User.Email", "Required, EmailAddress", "user_without_at_domain", "رفض: يرجى إدخال بريد إلكتروني صالح"),
        ("Hotel.Rating", "Range (1.0 إلى 5.0)", "إدخال القيمة 7.5", "رفض: التقييم يجب أن يكون بين 1 و 5"),
        ("Room.PricePerNight", "Range (1 إلى 1,000,000)", "إدخال قيمة سالبة (-200)", "رفض: يجب أن يكون السعر رقماً موجباً"),
        ("Room.Capacity", "Range (1 إلى 20 ضيفاً)", "إدخال القيمة 0", "رفض: سعة الغرفة يجب أن تتسع لشخص واحد على الأقل"),
        ("Booking.Dates", "IValidatableObject rule", "تاريخ المغادرة قبل تاريخ الوصول", "رفض: تاريخ المغادرة يجب أن يكون بعد الوصول")
    ]

    for row_idx, data in enumerate(val_data, start=1):
        row = table_val.rows[row_idx]
        bg = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 50, 50, 80, 80)
            p = cell.paragraphs[0]
            set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            r = p.add_run(val)
            r.font.name = "Segoe UI"
            r.font.size = Pt(8.5)
            set_run_rtl(r)

    add_rtl_heading(doc, "3.4 اختبار منع الحجز المزدوج (Double Booking Prevention Test):", level=2)
    add_rtl_paragraph(
        doc,
        "من أهم قواعد منطق الأعمال (Business Logic) المطبقة في النظام هي منع تضارب الحجوزات لنفس الغرفة. "
        "تم اختبار السيناريو التالي برمجياً:\n"
        "• تم حجز الغرفة رقم 101 من تاريخ 15 أكتوبر إلى 20 أكتوبر وتم تأكيد الحجز بنجاح.\n"
        "• تمت محاولة إرسال حجز آخر لنفس الغرفة من تاريخ 18 أكتوبر إلى 23 أكتوبر (فترة متداخلة).\n"
        "• النتيجة: رفض النظام الطلب فوراً برمز 409 Conflict مع ظهور رسالة واضحة: 'عذراً، هذه الغرفة محجوزة بالفعل في الفترة المحددة'، "
        "وظلت قاعدة البيانات في حالة سلامة تامة دون أي تكرار أو تعارض."
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # 6. الجزء الرابع: التوليف التراكمي لأسابيع المشروع (W2 - W9)
    # -------------------------------------------------------------
    add_rtl_heading(doc, "5. الجزء الرابع: التوليف التراكمي لمخرجات أسابيع المشروع (W2 إلى W9)", level=1)

    add_rtl_paragraph(
        doc,
        "يوضح هذا الجدول مطابقة إنجازات المشروع مع توزيع الدرجات والأسابيع المحددة في دليل مساق التدريب الميداني لجامعة الحكمة:"
    )

    table_weeks = doc.add_table(rows=10, cols=4)
    table_weeks.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_weeks, "E5E7EB", "4")

    cols_w = ["الأسبوع", "المتطلب الأكاديمي", "الوزن", "ما تم تنفيذه في المشروع"]
    for i, c in enumerate(cols_w):
        cell = table_weeks.rows[0].cells[i]
        set_cell_background(cell, "1E1B4B")
        p = cell.paragraphs[0]
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run(c)
        r.font.name = "Segoe UI"
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9)
        set_run_rtl(r)

    weeks_summary = [
        ("W2", "مقترح المشروع والمتطلبات (Project Proposal)", "5%", "صياغة نطاق العمل، دراسة المشكلة، و12 متطلباً وظيفياً و6 متطلبات غير وظيفية"),
        ("W3", "التصميم المعماري ومخططات UML (Clean Architecture)", "15%", "بناء معمارية Clean Architecture، ومخططات Use Case و Sequence و Class Diagrams"),
        ("W4", "تصميم واجهات وتجربة المستخدم (UI/UX Design)", "5%", "تصميم النماذج التفاعلية في Adobe XD بألوان Indigo و Violet ولمسات Gold الفاخرة"),
        ("W5", "تطوير واجهات البرمجة (Web API Development)", "10%", "بناء 18 نقطة نهاية برمجية مؤمنة بـ JWT مع حماية البيانات ومعالجة الأخطاء"),
        ("W6", "فحص واختبار واجهات البرمجة (API Testing)", "5%", "أتمتة الفحص عبر Postman Collection وتوفير واجهة Swagger UI التفاعلية"),
        ("W7", "تطبيق الويب المتكامل (ASP.NET Core MVC)", "25%", "بناء واجهات MVC بـ Bootstrap 5 RTL، خط Cairo، متغيرات CSS، و DataTables.js"),
        ("W8", "تطبيق الهاتف المحمول (Flutter Mobile App)", "15%", "بناء تطبيق موبايل متكامل بنمط Provider وقائمة Drawer و BottomNav و SheetModal"),
        ("W9", "إدارة الإصدارات والدمج (GitHub Workflow)", "5%", "تطبيق العمل على الفروع والدمج، وفصل تطبيق الموبايل في ريبو مستقل وفق طلب المشرف"),
        ("W10", "التوثيق النهائي والدليل الفني (Documentation)", "15%", "إعداد دليل التثبيت، دليل المستخدم، مصفوفة الاختبارات، الخاتمة، والمراجع العلمية")
    ]

    for row_idx, data in enumerate(weeks_summary, start=1):
        row = table_weeks.rows[row_idx]
        bg = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 50, 50, 80, 80)
            p = cell.paragraphs[0]
            set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            r = p.add_run(val)
            r.font.name = "Segoe UI"
            r.font.size = Pt(8.5)
            if col_idx == 2:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
            set_run_rtl(r)

    # -------------------------------------------------------------
    # 7. الجزء الخامس: الخاتمة والتحديات (W10.4 Conclusion)
    # -------------------------------------------------------------
    add_rtl_heading(doc, "6. الجزء الخامس: الخاتمة والتحديات الهندسية والتطويرات المستقبلية (W10.4)", level=1)

    add_rtl_paragraph(
        doc,
        "مثلت تجربة التدريب الميداني محاكاة متقدمة ومثمرة لبيئات الإنتاج الفعلية في كبرى الشركات الهندسية. "
        "تمكن المهندس من قيادة وتطوير كافة مراحل النظام بجهد فردي كامل، والتغلب على العديد من التحديات الهندسية المعقدة:"
    )

    challenges_ar = [
        ("فصل تطبيق الموبايل في مشروع ومستودع مستقل كلياً", "استجابة لتوجيه أ.د/ إبراهيم البلطة بضرورة استقلالية Flutter، تم استخراج المشروع بالكامل خارج حل الـ .NET، وضبط ملفات الـ gitignore لعزل ملفات البناء الثقيلة، وتهيئة مستودع مستقل على GitHub مع تاريخ دمج فروع كامل."),
        ("مزامنة التوقيتات القياسية العالمية (UTC Timestamps)", "تم التغلب على فروق التوقيت بين أجهزة العملاء وقاعدة البيانات من خلال توحيد كافة التواريخ بنظام DateTime.SpecifyKind(d, DateTimeKind.Utc) على مستوى طبقة الـ Application لمنع أي خطأ بيوم الحجز."),
        ("معالجة استجابة شاشات الموبايل المتفاوتة (Responsive Frame)", "تم ابتكار عنصر ResponsivePhoneFrame في Flutter لمنع ظهور أخطاء تجاوز البكسلات (Pixel Overflow) على الشاشات العريضة وضمان تجربة عرض مريحة ومتناسقة.")
    ]

    for title, desc in challenges_ar:
        p = add_rtl_paragraph(doc, "", space_after=3)
        r_t = p.add_run(f"• {title}: ")
        r_t.font.bold = True
        set_run_rtl(r_t)
        r_d = p.add_run(desc)
        set_run_rtl(r_d)

    add_rtl_heading(doc, "خارطة التطويرات المستقبلية للنظام (Future Enhancements):", level=2)
    add_rtl_paragraph(doc, "1. ربط بوابات الدفع الإلكتروني الدولية والمحلية (Stripe / PayPal / المحافظ الإلكترونية اليمنية).")
    add_rtl_paragraph(doc, "2. تفعيل الإشعارات الفورية السحابية (Push Notifications) عبر Firebase Cloud Messaging.")
    add_rtl_paragraph(doc, "3. دمج محرك توصيات ذكي بالذكاء الاصطناعي لاقتراح الغرف والفنادق بناءً على تفضيلات النزلاء السابقة.")

    # -------------------------------------------------------------
    # 8. الجزء السادس: المراجع العلمية (W10.5 References)
    # -------------------------------------------------------------
    add_rtl_heading(doc, "7. الجزء السادس: المراجع الأكاديمية والتقنية المعتمدة (W10.5 References)", level=1)

    refs_ar = [
        "روبرت سي مارتن (Uncle Bob) — كتاب Clean Architecture: A Craftsman's Guide to Software Structure and Design، دار النشر Prentice Hall.",
        "شركة مايكروسوفت العالمية — الوثائق الرسمية والمراجع الهندسية لإطار العمل ASP.NET Core 10.0 و Entity Framework Core، منصة Microsoft Learn.",
        "شركة جوجل — التوثيق الرسمي ومستندات المطورين لإطار العمل Flutter وإدارة الحالة عبر Provider Pattern، منصة Flutter.dev.",
        "د. روي فيلدينغ — أطروحة المعمارية البرمجية لشبكات الويب وتصميم برمجيات الـ REST APIs، جامعة كاليفورنيا.",
        "فريق مهندسي الإنترنت (IETF) — المعيار القياسي الدولي RFC 7519 لرموز التوثيق المشفرة JSON Web Token (JWT).",
        "جامعة الحكمة — دليل مساق التدريب الميداني لمشاريع تطوير البرمجيات الشاملة الفردية (SDLC)، قسم تكنولوجيا المعلومات، صنعاء، الجمهورية اليمنية."
    ]

    for idx, ref in enumerate(refs_ar, start=1):
        p = add_rtl_paragraph(doc, f"[{idx}] {ref}", space_after=3)

    # حفظ المستند في المكانين
    output_dir_1 = r"E:\Project_ASP.NET"
    output_path_1 = os.path.join(output_dir_1, "Hotel_Booking_System_W10_Documentation_AR.docx")
    doc.save(output_path_1)
    print("Arabic document saved to:", output_path_1)

    output_dir_2 = r"E:\Project_ASP.NET\HOTEL_BOOKIN_SYSTEM\Hotel_Booking_System\docs"
    output_path_2 = os.path.join(output_dir_2, "Hotel_Booking_System_W10_Documentation_AR.docx")
    doc.save(output_path_2)
    print("Arabic document copy saved to:", output_path_2)

if __name__ == "__main__":
    build_arabic_document()
