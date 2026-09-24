import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

FONT_ARABIC = "Traditional Arabic"
FONT_ENGLISH = "Arial"
FONT_CODE = "Consolas"

# Official Academic & Corporate Color Palette
COLOR_PRIMARY_DARK = RGBColor(0x1E, 0x1B, 0x4B)   # كحلي إمبراطوري داكن للعناوين الرئيسية
COLOR_BRAND = RGBColor(0x37, 0x30, 0xA3)          # نيلي ملكي للعناوين الفرعية
COLOR_ACCENT = RGBColor(0xC9, 0x97, 0x1F)         # ذهبي فاخر للأطر والشارات الأكاديمية
COLOR_TEXT = RGBColor(0x0F, 0x17, 0x2A)           # كحلي فحمي داكن عريض وواضح للقراءة
COLOR_MUTED = RGBColor(0x33, 0x41, 0x55)          # رمادي ناصع وواضح
COLOR_GREEN = RGBColor(0x06, 0x5F, 0x46)          # أخضر رسمي داكن لنتائج الفحص الناجحة
COLOR_AMBER = RGBColor(0x92, 0x40, 0x0E)          # بني كهرماني للتنبيهات

def set_paragraph_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    """Applies true bidirectional RTL and justification to paragraph."""
    pPr = p._p.get_or_add_pPr()
    bidi = parse_xml(f'<w:bidi {nsdecls("w")}/>')
    pPr.append(bidi)
    p.alignment = align
    if align == WD_ALIGN_PARAGRAPH.RIGHT:
        jc = parse_xml(f'<w:jc {nsdecls("w")} w:val="both"/>')
        pPr.append(jc)

def set_run_font(run, font_name=FONT_ARABIC, size_pt=14, bold=False, italic=False, color=COLOR_TEXT):
    """Sets explicit font properties for both Complex Scripts (Arabic) and Latin."""
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(
        f'<w:rFonts {nsdecls("w")} '
        f'w:ascii="{FONT_ENGLISH}" '
        f'w:hAnsi="{FONT_ENGLISH}" '
        f'w:cs="{font_name}"/>'
    )
    rtl = parse_xml(f'<w:rtl {nsdecls("w")}/>')
    rPr.append(rFonts)
    rPr.append(rtl)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_styling(table, border_color="CBD5E1", border_sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>'
        f'<w:insideH w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    bidiVisual = parse_xml(f'<w:bidiVisual {nsdecls("w")}/>')
    tblPr.append(bidiVisual)

def add_arabic_paragraph(doc, text="", space_after=6, bold=False, size=14, color=COLOR_TEXT, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = doc.add_paragraph()
    set_paragraph_rtl(p, align)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    if text:
        r = p.add_run(text)
        set_run_font(r, font_name=FONT_ARABIC, size_pt=size, bold=bold, color=color)
    return p

def add_arabic_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
    p.paragraph_format.keep_with_next = True
    
    if level == 1:
        p.paragraph_format.space_before = Pt(22)
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(text)
        set_run_font(r, font_name=FONT_ARABIC, size_pt=18.5, bold=True, color=COLOR_PRIMARY_DARK)
        pBdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="18" w:space="6" w:color="3730A3"/>'
            f'</w:pBdr>'
        )
        p._p.get_or_add_pPr().append(pBdr)
    elif level == 2:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        set_run_font(r, font_name=FONT_ARABIC, size_pt=16, bold=True, color=COLOR_BRAND)
    else:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        set_run_font(r, font_name=FONT_ARABIC, size_pt=14.5, bold=True, color=COLOR_PRIMARY_DARK)
    return p

def add_card_box(doc, title, content_list, border_color="3730A3", bg_color="EEF2FF"):
    """Adds a stylish callout card box with a solid right border and soft background."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="single" w:sz="32" w:space="0" w:color="{border_color}"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    tblPr.append(parse_xml(f'<w:bidiVisual {nsdecls("w")}/>'))
    
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, 120, 120, 180, 180)
    
    p0 = cell.paragraphs[0]
    set_paragraph_rtl(p0, WD_ALIGN_PARAGRAPH.RIGHT)
    p0.paragraph_format.space_after = Pt(4)
    r_t = p0.add_run(title)
    set_run_font(r_t, font_name=FONT_ARABIC, size_pt=15, bold=True, color=COLOR_PRIMARY_DARK)
    
    for item in content_list:
        p = cell.add_paragraph()
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(item)
        set_run_font(r, font_name=FONT_ARABIC, size_pt=13.5, bold=False, color=COLOR_TEXT)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="1E1B4B"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, 100, 100, 140, 140)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(code_text)
    r.font.name = FONT_CODE
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_figure_with_caption(doc, image_path, caption_arabic, width_inches=6.0, is_vertical=False):
    """Embeds an actual project screenshot inside an academic figure card with caption."""
    if not os.path.exists(image_path):
        print(f"Warning: Image not found: {image_path}")
        return

    tbl = doc.add_table(rows=2, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    tblPr.append(parse_xml(f'<w:bidiVisual {nsdecls("w")}/>'))

    # Row 0: Screenshot image
    cell_img = tbl.rows[0].cells[0]
    cell_img.width = Inches(width_inches)
    set_cell_background(cell_img, "FFFFFF")
    set_cell_margins(cell_img, top=100, bottom=60, left=100, right=100)
    p_img = cell_img.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_after = Pt(2)
    r_img = p_img.add_run()
    
    img_w = width_inches - 0.4 if not is_vertical else 3.4
    r_img.add_picture(image_path, width=Inches(img_w))

    # Row 1: Figure Caption
    cell_cap = tbl.rows[1].cells[0]
    set_cell_background(cell_cap, "F1F5F9")
    set_cell_margins(cell_cap, top=60, bottom=80, left=120, right=120)
    p_cap = cell_cap.paragraphs[0]
    set_paragraph_rtl(p_cap, WD_ALIGN_PARAGRAPH.CENTER)
    p_cap.paragraph_format.space_after = Pt(0)
    r_cap = p_cap.add_run(caption_arabic)
    set_run_font(r_cap, font_name=FONT_ARABIC, size_pt=12.5, bold=True, color=COLOR_PRIMARY_DARK)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def build_curated_w10_document():
    doc = docx.Document()
    sc_dir = r"E:\Project_ASP.NET\HOTEL_BOOKIN_SYSTEM\Hotel_Booking_System\docs\screenshots"

    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
        sectPr = section._sectPr
        bidi = parse_xml(f'<w:bidi {nsdecls("w")}/>')
        sectPr.append(bidi)

    # -------------------------------------------------------------
    # 1. صفحة الغلاف الأكاديمية الرسمية
    # -------------------------------------------------------------
    p_inst = doc.add_paragraph()
    set_paragraph_rtl(p_inst, WD_ALIGN_PARAGRAPH.CENTER)
    p_inst.paragraph_format.space_after = Pt(20)
    r_inst = p_inst.add_run(
        "الجمهورية اليمنية\n"
        "وزارة التعليم العالي والبحث العلمي\n"
        "جامعة الحكمة — كلية العلوم والتقنية\n"
        "قسم تقنية المعلومات (Information Technology)"
    )
    set_run_font(r_inst, font_name=FONT_ARABIC, size_pt=14.5, bold=True, color=COLOR_MUTED)

    # الخط الفاصل الذهبي
    p_div = doc.add_paragraph()
    set_paragraph_rtl(p_div, WD_ALIGN_PARAGRAPH.CENTER)
    p_div.paragraph_format.space_after = Pt(25)
    r_div = p_div.add_run("❖  ❖  ❖")
    set_run_font(r_div, font_name=FONT_ARABIC, size_pt=15, bold=True, color=COLOR_ACCENT)

    # عنوان الوثيقة الأكاديمي
    p_title = doc.add_paragraph()
    set_paragraph_rtl(p_title, WD_ALIGN_PARAGRAPH.CENTER)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("التقرير الفني والتوثيق الأكاديمي الشامل للمشروع النهائي")
    set_run_font(r_title, font_name=FONT_ARABIC, size_pt=23, bold=True, color=COLOR_PRIMARY_DARK)

    # العنوان الفرعي
    p_sub = doc.add_paragraph()
    set_paragraph_rtl(p_sub, WD_ALIGN_PARAGRAPH.CENTER)
    p_sub.paragraph_format.space_after = Pt(20)
    r_sub = p_sub.add_run(
        "منظومة إدارة وحجز الفنادق المتكاملة والموزعة\n"
        "Hotel Booking Management System (Enterprise Architecture)\n"
        "مخرجات الأسبوع العاشر: التثبيت، دليل المستخدم، مصفوفة الاختبارات، الخاتمة، والمراجع"
    )
    set_run_font(r_sub, font_name=FONT_ARABIC, size_pt=15, bold=True, color=COLOR_BRAND)

    # شارة مساق التدريب الميداني
    badge_tbl = doc.add_table(rows=1, cols=1)
    badge_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    b_cell = badge_tbl.rows[0].cells[0]
    b_cell.width = Inches(5.8)
    set_cell_background(b_cell, "EEF2FF")
    set_cell_margins(b_cell, 100, 100, 150, 150)
    p_badge = b_cell.paragraphs[0]
    set_paragraph_rtl(p_badge, WD_ALIGN_PARAGRAPH.CENTER)
    r_badge = p_badge.add_run("مشروع تسليم متطلبات مساق التدريب الميداني الفردي (Field Training Course - W10)")
    set_run_font(r_badge, font_name=FONT_ARABIC, size_pt=13.5, bold=True, color=COLOR_PRIMARY_DARK)

    doc.add_paragraph().paragraph_format.space_after = Pt(30)

    # جدول بيانات الطالب والإشراف الأكاديمي
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_styling(meta_table, border_color="CBD5E1", border_sz="6")

    meta_items = [
        ("إعداد مهندس البرمجيات / الطالب:", "لؤي الصروري (Loai Alsorory)"),
        ("المشرف الأكاديمي على المساق:", "أ.د / إبراهيم أحمد البلطة"),
        ("الدرجة العلمية المستهدفة:", "بكالوريوس في تقنية المعلومات (B.Sc. in IT)"),
        ("طبيعة العمل والتنفيذ:", "تطوير فردي شامل (Individual Full-Stack SDLC)"),
        ("العام الجامعي وموسم التسليم:", "2024 - 2025م  |  الفصل الدراسي الثاني")
    ]

    for idx, (label, val) in enumerate(meta_items):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.6)
        cell_val.width = Inches(3.9)
        set_cell_background(cell_lbl, "F8FAFC")
        set_cell_margins(cell_lbl, 90, 90, 130, 130)
        set_cell_margins(cell_val, 90, 90, 130, 130)

        p0 = cell_lbl.paragraphs[0]
        set_paragraph_rtl(p0, WD_ALIGN_PARAGRAPH.RIGHT)
        r0 = p0.add_run(label)
        set_run_font(r0, font_name=FONT_ARABIC, size_pt=13, bold=True, color=COLOR_PRIMARY_DARK)

        p1 = cell_val.paragraphs[0]
        set_paragraph_rtl(p1, WD_ALIGN_PARAGRAPH.RIGHT)
        r1 = p1.add_run(val)
        set_run_font(r1, font_name=FONT_ARABIC, size_pt=13, bold=False, color=COLOR_TEXT)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 2. المقدمة والملخص التنفيذي والمعمارية
    # -------------------------------------------------------------
    add_arabic_heading(doc, "1. المقدمة والملخص التنفيذي والمعمارية البرمجية للمشروع", level=1)
    
    add_arabic_paragraph(
        doc,
        "يُعد هذا التقرير التوثيقي المرجع الفني والأكاديمي الشامل لتسليم متطلبات مساق التدريب الميداني بجامعة الحكمة، "
        "المُقام تحت إشراف أ.د/ إبراهيم أحمد البلطة. تم تصميم وتنفيذ المشروع بمحاكاة دقيقة لبيئات شركات البرمجيات العالمية، "
        "حيث تولى الطالب دور مهندس برمجيات متكامل (Full-Stack Software Engineer) لتنفيذ كافة مراحل دورة حياة تطوير النظم (SDLC) "
        "بشكل فردي وصارم، انطلاقاً من التحليل والنمذجة والتصميم المعماري، ومروراً بتطوير قواعد البيانات والواجهات البرمجية والتطبيقات، "
        "وصولاً إلى الفحص وضمان الجودة والتوثيق المتكامل."
    )
    
    add_arabic_paragraph(
        doc,
        "يقدم النظام حلاً هندسياً شاملاً وموزعاً لإدارة قطاع الضيافة وحجوزات الفنادق، مبنياً وفق أحدث المعايير البرمجية "
        "ويتوزع على مشروعين ومستودعين منفصلين تماماً تنفيذاً لقرار العزل المعماري الصادر عن المشرف الأكاديمي:"
    )

    add_card_box(
        doc,
        "المشروع الأول: الباك إند وبوابة الويب الإدارية (HotelBookingSystem)",
        [
            "• إطار العمل والتقنيات: مبني بتقنية ASP.NET Core 10 وقاعدة بيانات SQL Server عبر Entity Framework Core 10.",
            "• النمط المعماري: معمارية البصلة / النظيفة (Clean Architecture) مع نمط المستودع (Repository Pattern) لضمان استقلالية طبقات العمل.",
            "• المكونات: يضم واجهات برمجة التطبيقات (RESTful Web APIs) وبوابة إدارة متكاملة بنظام ASP.NET Core MVC (Razor Views) مدعومة بـ Bootstrap 5 RTL وخط Cairo و DataTables.js."
        ],
        border_color="3730A3",
        bg_color="EEF2FF"
    )

    add_card_box(
        doc,
        "المشروع الثاني: تطبيق الهاتف المحمول المستقل (HotelBookingSystem-Flutter)",
        [
            "• إطار العمل والتقنيات: مبني بتقنية Google Flutter مع لغة Dart 3.x ونمط إدارة الحالة المتقدم (Provider Pattern).",
            "• الفصل المعماري: تم فصله في مشروع ومستودع مستقل كلياً على GitHub كعميل مستقل (Standalone Micro-Client) استجابة لتوجيه د/ إبراهيم البلطة.",
            "• التكامل: يتصل بالباك إند حصرياً عبر بروتوكول HTTP/JSON، مع دعم واجهات عربية أصيلة (RTL)، وقائمة جانبية Drawer، وشريط تنقل سفلي BottomNav، ونافذة حجز منبثقة SheetModal."
        ],
        border_color="7C3AED",
        bg_color="FAF5FF"
    )

    add_arabic_heading(doc, "جدول تفصيل الطبقات المعمارية للمشروع (Clean Architecture Layers):", level=2)

    table_arch = doc.add_table(rows=6, cols=3)
    table_arch.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_styling(table_arch, border_color="CBD5E1", border_sz="4")

    headers_arch = ["الطبقة البرمجية (Layer)", "المسؤولية والوظيفة في النظام", "التقنيات والمكتبات المستخدمة"]
    for i, h in enumerate(headers_arch):
        cell = table_arch.rows[0].cells[i]
        set_cell_background(cell, "1E1B4B")
        p = cell.paragraphs[0]
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run(h)
        set_run_font(r, font_name=FONT_ARABIC, size_pt=13, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    arch_data = [
        ("HotelBookingSystem.Domain", "تحتوي الكيانات الأساسية (Hotels, Rooms, Bookings, Users)، أنواع البيانات، وقواعد العمل المجردة بدون أي اعتمادية على مكتبات خارجية.", "C# 13, Plain Old CLR Objects (POCOs)"),
        ("HotelBookingSystem.Application", "تحتوي كائنات نقل البيانات (DTOs)، واجهات المستودعات (IRepositories)، خدمات الأعمال، وضوابط التحقق Validation.", "FluentValidation, DataAnnotations"),
        ("HotelBookingSystem.Infrastructure", "تحتوي سياق قاعدة البيانات (ApplicationDbContext)، وتطبيق المستودعات عبر EF Core، وإدارة الترحيلات ومحرك زرع البيانات DbSeeder.", "Entity Framework Core 10, SQL Server"),
        ("HotelBookingSystem.API", "المحطة الطرفية لتوفير الـ Endpoints للعملاء الخارجيين، وتوثيق Swagger، وتأمين المصادقة عبر JWT Bearer Tokens.", "ASP.NET Core 10 Web API, Swagger/OpenAPI"),
        ("HotelBookingSystem.Web", "بوابة الإدارة المركزية بنمط MVC، تضم شاشات التحكم، وعرض الإحصائيات، ونماذج التعديل والحذف مع نوافذ Bootstrap Modals.", "ASP.NET Core MVC, Razor, Bootstrap 5 RTL, DataTable.js")
    ]

    for row_idx, data in enumerate(arch_data, start=1):
        row = table_arch.rows[row_idx]
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 110, 110)
            p = cell.paragraphs[0]
            set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            r = p.add_run(text)
            set_run_font(r, font_name=FONT_ARABIC, size_pt=12.5, bold=(col_idx == 0), color=COLOR_TEXT)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 3. الجزء الأول: دليل التثبيت والتجهيز البيئي (W10.1)
    # -------------------------------------------------------------
    add_arabic_heading(doc, "2. الجزء الأول: دليل التثبيت والتجهيز البيئي والتشغيل (W10.1 Installation Guide)", level=1)

    add_arabic_paragraph(
        doc,
        "يحدد هذا القسم المتطلبات الهندسية والخطوات الإجرائية الدقيقة لتجهيز بيئة العمل، وإعداد قواعد البيانات، "
        "وترجمة وتشغيل كافة مكونات النظام بنجاح من المصدر (Source Code)."
    )

    add_arabic_heading(doc, "1.2 المتطلبات المسبقة للبيئة والعتاد (Prerequisites):", level=2)

    table_req = doc.add_table(rows=6, cols=3)
    table_req.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_styling(table_req, border_color="CBD5E1", border_sz="4")

    headers_req = ["المكون البرمجي / التقنية", "الحد الأدنى المطلوب للتشغيل", "الغرض والدور في النظام"]
    for i, h in enumerate(headers_req):
        cell = table_req.rows[0].cells[i]
        set_cell_background(cell, "1E1B4B")
        p = cell.paragraphs[0]
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run(h)
        set_run_font(r, font_name=FONT_ARABIC, size_pt=13, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    req_data = [
        (".NET SDK", "الإصدار .NET 10.0 SDK (x64)", "بيئة بناء وتشغيل مشروع الباك إند API وبوابة الويب MVC"),
        ("محرك قاعدة البيانات", "SQL Server LocalDB أو SQL Server 2022", "مستودع حفظ البيانات العلائقية للفنادق والغرف والحجوزات والمستخدمين"),
        ("حزمة أدوات الموبايل", "Flutter SDK 3.19 فما فوق مع Dart 3", "بناء وتشغيل تطبيق الهاتف المحمول عبر المتصفح أو أجهزة أندرويد"),
        ("متصفح الويب الحديث", "Google Chrome 120 فما فوق", "استعراض بوابة الويب وتوثيق Swagger واختبار تطبيق Flutter Web"),
        ("نظام إدارة الإصدارات", "Git CLI 2.40 فما فوق", "مزامنة التحديثات، إدارة الفروع، ورفع المشاريع لمستودعات GitHub")
    ]

    for row_idx, data in enumerate(req_data, start=1):
        row = table_req.rows[row_idx]
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 110, 110)
            p = cell.paragraphs[0]
            set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            r = p.add_run(text)
            set_run_font(r, font_name=FONT_ARABIC, size_pt=12.5, bold=(col_idx == 0), color=COLOR_TEXT)

    add_arabic_heading(doc, "2.2 إعداد قاعدة البيانات والزرع التلقائي للبيانات (Automated Seeding Engine):", level=2)
    add_arabic_paragraph(
        doc,
        "تم تضمين منطق تهيئة ذاتي ومتقدم لقاعدة البيانات داخل ملف Program.cs؛ حيث يتبع النظام استراتيجية "
        "الإنشاء الذاتي (Self-Bootstrapping Database) التي تُغني المشرف أو المقيّم عن كتابة أي أوامر SQL يدوية:\n"
        "1. التحقق التلقائي: عند تشغيل مشروع الـ Web أو الـ API، يستدعي النظام Database.EnsureCreated() الذي يتحقق من وجود قاعدة بيانات HotelBookingDb، وفي حال عدم وجودها يتم إنشاؤها فوراً وبناء الجداول الأربعة والعلاقات والمفاتيح الأجنبية وفهارس البحث.\n"
        "2. محرك الزرع الذكي (DbSeeder.SeedAsync): يتحقق النظام هل الجداول فارغة، فإذا كانت فارغة يقوم بملئها فورياً ببيانات واقعية وفخمة تشمل:\n"
        "   - 12 فندقاً فريداً: موزعة على كبرى العواصم والمدن العربية (الرياض، جدة، دبي، الدوحة، مسقط، الكويت، القاهرة، بيروت، وغيرها)، ولكل فندق صورة فوتوغرافية معمارية فريدة وعالية الجودة.\n"
        "   - 48 غرفة فندقية مستقلة: بمعدل 4 غرف لكل فندق تتنوع بين (مفردة، مزدوجة، أجنحة عائلية، وديلوكس فاخر)، بأسعار مدروسة وسعات استيعابية من 1 إلى 10 ضيوف، وتم تخصيص صورة غرفة فريدة ومستقلة لكل غرفة تختلف تماماً عن صور الفنادق.\n"
        "   - بيانات الاعتماد المشفرة: إنشاء حساب مدير النظام (admin@hotel.com) وحسابات عملاء تجريبية بكلمات مرور محمية بخوارزمية التشفير PBKDF2 والـ Salt العشوائي."
    )

    add_arabic_heading(doc, "3.2 خطوات وأوامر تشغيل المنظومة خطوة بخطوة (Execution Guide):", level=2)

    add_arabic_paragraph(doc, "أولاً: استنساخ وبناء مشروع الباك إند والويب (ASP.NET Core Solution):", bold=True, size=13.5, color=COLOR_BRAND)
    add_code_block(
        doc,
        "git clone https://github.com/LoaiAlsorory/HotelBookingSystem.git\n"
        "cd HotelBookingSystem\n"
        "dotnet restore HotelBookingSystem.sln\n"
        "dotnet build HotelBookingSystem.sln --configuration Release --nologo"
    )

    add_arabic_paragraph(doc, "ثانياً: تشغيل خدمة واجهات البرمجة (RESTful Web API):", bold=True, size=13.5, color=COLOR_BRAND)
    add_code_block(
        doc,
        "cd src/HotelBookingSystem.API\n"
        "dotnet run\n"
        "# السيرفر يعمل على: http://localhost:5242\n"
        "# واجهة توثيق Swagger التفاعلية: http://localhost:5242/swagger"
    )

    # صورة 1 المختارة: واجهة Swagger
    add_figure_with_caption(
        doc,
        os.path.join(sc_dir, "fig07_api_swagger.png"),
        "الشكل (1): واجهة توثيق واختبار نقاط النهاية البرمجية (RESTful Web API) عبر Swagger OpenAPI"
    )

    add_arabic_paragraph(doc, "ثالثاً: تشغيل بوابة الويب الإدارية (ASP.NET Core MVC):", bold=True, size=13.5, color=COLOR_BRAND)
    add_code_block(
        doc,
        "cd src/HotelBookingSystem.Web\n"
        "dotnet run\n"
        "# رابط البوابة الإدارية: http://localhost:5200\n"
        "# حساب تسجيل دخول المدير: admin@hotel.com  |  كلمة المرور: Password123!"
    )

    add_arabic_paragraph(doc, "رابعاً: تشغيل تطبيق الهاتف المحمول المستقل (Flutter Mobile Client):", bold=True, size=13.5, color=COLOR_BRAND)
    add_code_block(
        doc,
        "cd E:/Project_ASP.NET/HOTEL_BOOKIN_SYSTEM/Hotel_Booking_Flutter\n"
        "flutter pub get\n"
        "flutter run -d chrome     # للتشغيل المباشر عبر المتصفح\n"
        "flutter run              # للتشغيل على هاتف أندرويد متصل عبر USB"
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # 4. الجزء الثاني: دليل المستخدم المصور (W10.2 User Manual)
    # -------------------------------------------------------------
    add_arabic_heading(doc, "3. الجزء الثاني: دليل المستخدم الشامل وواجهات النظام (W10.2 User Manual)", level=1)

    add_arabic_paragraph(
        doc,
        "صُمم النظام لخدمة شريحتين أساسيتين من المستخدمين عبر واجهتين هندسيتين متطورتين؛ بوابة الإدارة المكتبية (Web Portal) "
        "التي تلبي احتياجات مسؤولي النظام ومدراء الفنادق، وتطبيق الهاتف المحمول (Flutter App) المصمم لتوفير تجربة حجز عصرية وسريعة للنزلاء. "
        "وفيما يلي استعراض لأهم الواجهات الجوهرية المدعومة بالصور التوضيحية المنتقاة بعناية وبطاقات الشرح الفني:"
    )

    add_arabic_heading(doc, "1.3 دليل استخدام بوابة الويب الإدارية (ASP.NET Core MVC Portal):", level=2)

    # تفاصيل تسجيل الدخول
    add_card_box(
        doc,
        "أ. واجهة تسجيل الدخول والمصادقة (Split-Screen Authentication)",
        [
            "تعتمد واجهة تسجيل الدخول تصميماً احترافياً منقسماً يجمع بين جمالية الصورة الترحيبية ونموذج الإدخال الآمن، "
            "حيث يتم التحقق من بيانات الدخول ضد التشفير المعتمد وإصدار كوكيز مصادقة آمنة لجلسة العمل."
        ],
        border_color="3730A3",
        bg_color="EEF2FF"
    )

    # صورة 2 المختارة: لوحة التحكم
    add_card_box(
        doc,
        "ب. لوحة التحكم المركزية وبوابة الإحصائيات (Executive Dashboard)",
        [
            "تُمثل لوحة التحكم القلب النابض لإدارة المنظومة، حيث تستعرض 4 بطاقات إحصائية حية (KPIs) تعرض إجمالي الفنادق، "
            "والغرف الشاغرة، والحجوزات المؤكدة، ومجموع الإيرادات، مع جدول تفاعلي لأحدث الحجوزات وأزرار إجراءات سريعة."
        ],
        border_color="3730A3",
        bg_color="EEF2FF"
    )
    add_figure_with_caption(
        doc,
        os.path.join(sc_dir, "fig02_web_dashboard.png"),
        "الشكل (2): لوحة التحكم المركزية وبطاقات مؤشرات الأداء الحية وجدول الحجوزات (Main Dashboard)"
    )

    # صورة 3 المختارة: إدارة الغرف
    add_card_box(
        doc,
        "ج. شاشة إدارة الغرف الفندقية والمخزون وحالات التوفر (Rooms Inventory)",
        [
            "تتيح للمشرفين فحص مخزون الغرف، ومتابعة أسعار الإقامة وتصنيفاتها (مفردة، مزدوجة، أجنحة، ديلوكس)، "
            "مع إمكانية تعديل حالة الغرفة بين (متاحة ومحجوزة) وتحديث الطاقة الاستيعابية بضغطة زر واحدة."
        ],
        border_color="3730A3",
        bg_color="EEF2FF"
    )
    add_figure_with_caption(
        doc,
        os.path.join(sc_dir, "fig04_web_rooms.png"),
        "الشكل (3): شاشة إدارة الغرف الفندقية وتحديد الأسعار وحالات التوفر (Rooms Inventory)"
    )

    # صورة 4 المختارة: إدارة الحجوزات
    add_card_box(
        doc,
        "د. شاشة إدارة ومتابعة دورة حياة الحجوزات والعمليات المالية (Bookings Workflow)",
        [
            "توفر رؤية شاملة لكافة الحجوزات، مع تواريخ الوصول والمغادرة واحتساب القيمة المالية تلقائياً بضرب عدد الليالي "
            "في سعر الغرفة، مع تفعيل محرك منع التضارب الزمني، وإمكانية تعديل حالة الحجز بين (مؤكد، معلق، وملغي)."
        ],
        border_color="3730A3",
        bg_color="EEF2FF"
    )
    add_figure_with_caption(
        doc,
        os.path.join(sc_dir, "fig05_web_bookings.png"),
        "الشكل (4): شاشة إدارة ومتابعة دورة حياة الحجوزات والعمليات المالية (Bookings Management)"
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # دليل تطبيق الموبايل مدعماً بالصور المنتقاة
    # -------------------------------------------------------------
    add_arabic_heading(doc, "2.3 دليل استخدام تطبيق الهاتف المحمول (Flutter Mobile Client):", level=2)

    add_arabic_paragraph(
        doc,
        "يتميز تطبيق الهاتف المحمول باستقلاليته الكاملة وسرعة استجابته وتوافقه التام مع واجهات الاستخدام العربية (RTL). "
        "تم تصميمه بنمط إدارة الحالة Provider، ويتضمن تجربة مستخدم سلسة وقائمة منسحبة وشريط تنقل سفلي."
    )

    # صورة 5 المختارة: الشاشة الرئيسية للموبايل
    add_card_box(
        doc,
        "أ. الواجهة الرئيسية وتصفح الفنادق (Mobile HomeScreen & Featured Hotels)",
        [
            "تستعرض الواجهة الرئيسية ترحيباً مخصصاً باسم النزيل، وبطاقات إحصائية، وشريطاً أفقياً فاخراً للفنادق المميزة "
            "مع تقييماتها، مما يوفر وصولاً سريعاً ومباشراً لأفضل خيارات الإقامة المتاحة."
        ],
        border_color="7C3AED",
        bg_color="FAF5FF"
    )
    add_figure_with_caption(
        doc,
        os.path.join(sc_dir, "fig08_mobile_dashboard.png"),
        "الشكل (5): الشاشة الرئيسية لتطبيق الهاتف المحمول وتصفح الفنادق الموصى بها (HomeScreen Dashboard)"
    )

    # صورة 6 المختارة: النافذة السفلية SheetModal
    add_card_box(
        doc,
        "ب. نافذة حجز الغرف التفاعلية السفلية (Interactive SheetModal Booking)",
        [
            "عند رغبة النزيل في حجز أي غرفة، تنبثق نافذة سفلية أنيقة وسلسة (BottomSheet) تتيح له تحديد تاريخ الدخول والمغادرة، "
            "واختيار عدد النزلاء، وتستعرض السعر الإجمالي المحتسب آلياً وتفاصيل العملية مع زر تأكيد الحجز الفوري."
        ],
        border_color="7C3AED",
        bg_color="FAF5FF"
    )
    add_figure_with_caption(
        doc,
        os.path.join(sc_dir, "fig12_mobile_sheet_modal.png"),
        "الشكل (6): النافذة السفلية المنبثقة لحجز الغرفة واحتساب الإجمالي تلقائياً (Interactive SheetModal)"
    )

    add_arabic_heading(doc, "جدول ملخص الخصائص والميزات التشغيلية لواجهات النظام:", level=2)

    table_features = doc.add_table(rows=7, cols=3)
    table_features.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_styling(table_features, border_color="CBD5E1", border_sz="4")

    headers_feat = ["الواجهة / المكون البرمجي", "المنصة المستهدفة", "الوظائف والخصائص الهندسية المتاحة"]
    for i, h in enumerate(headers_feat):
        cell = table_features.rows[0].cells[i]
        set_cell_background(cell, "1E1B4B")
        p = cell.paragraphs[0]
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run(h)
        set_run_font(r, font_name=FONT_ARABIC, size_pt=13, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    feat_data = [
        ("بوابة الدخول والمصادقة", "الويب (MVC)", "تسجيل دخول آمن مشفر بـ Cookies Authentication مع التحقق من الصلاحيات"),
        ("لوحة التحكم المركزية", "الويب (MVC)", "عرض 4 بطاقات إحصائية حية لـ KPIs وجدول أحدث العمليات وأزرار الإجراءات"),
        ("إدارة الفنادق والغرف", "الويب (MVC)", "جداول تفاعلية مع DataTable.js تدعم البحث والفرز الفوري ونوافذ تأكيد الحذف"),
        ("إدارة ومتابعة الحجوزات", "الويب (MVC)", "تسجيل وتعديل الحجوزات وتغيير الحالات ومنع التضارب واحتساب المبالغ آلياً"),
        ("الشاشة الرئيسية والقائمة", "الموبايل (Flutter)", "قائمة جانبية RTL Drawer، شريط سفلي BottomNav، وبطاقات الفنادق المميزة"),
        ("نافذة الحجز المنبثقة", "الموبايل (Flutter)", "حجز فوري بنمط ModalBottomSheet مع اختيار التواريخ واحتساب الإجمالي")
    ]

    for row_idx, data in enumerate(feat_data, start=1):
        row = table_features.rows[row_idx]
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 110, 110)
            p = cell.paragraphs[0]
            set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            r = p.add_run(text)
            set_run_font(r, font_name=FONT_ARABIC, size_pt=12.5, bold=(col_idx == 0), color=COLOR_TEXT)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 5. الجزء الثالث: تقرير الفحص وضمان الجودة (W10.3 Testing)
    # -------------------------------------------------------------
    add_arabic_heading(doc, "4. الجزء الثالث: تقرير الفحص وضمان الجودة والاختبارات (W10.3 Testing & QA)", level=1)

    add_arabic_paragraph(
        doc,
        "خضعت المنظومة البرمجية لمصفوفة اختبارات صارمة ومتعددة المستويات شملت اختبارات واجهات البرمجة (API Testing)، "
        "واختبارات التحقق من سلامة المدخلات (Input Validation)، واختبارات منطق الأعمال وحماية العمليات المتزامنة (Concurrency & Business Rules)."
    )

    add_arabic_heading(doc, "1.4 مصفوفة نتائج اختبارات واجهات الـ RESTful Web API عبر Postman و Swagger:", level=2)

    table_tests = doc.add_table(rows=11, cols=5)
    table_tests.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_styling(table_tests, border_color="CBD5E1", border_sz="4")

    cols_test = ["نقطة النهاية (Endpoint)", "نوع الطلب", "سيناريو الاختبار والفحص", "رمز الاستجابة المتوقع", "النتيجة الفعلية"]
    for i, c in enumerate(cols_test):
        cell = table_tests.rows[0].cells[i]
        set_cell_background(cell, "1E1B4B")
        p = cell.paragraphs[0]
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run(c)
        set_run_font(r, font_name=FONT_ARABIC, size_pt=12.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    tests_data = [
        ("/api/auth/login", "POST", "تسجيل دخول صحيح لمدير النظام", "200 OK", "ناجح (استلام JWT Token)"),
        ("/api/auth/login", "POST", "محاولة دخول بكلمة مرور خاطئة", "401 Unauthorized", "ناجح (رفض وتنبيه بالخطأ)"),
        ("/api/hotels", "GET", "استرجاع قائمة الفنادق كاملة", "200 OK", "ناجح (استرجاع 12 فندقاً)"),
        ("/api/hotels/{id}", "GET", "استرجاع فندق برقم المعرف 1", "200 OK", "ناجح (استرجاع بيانات الفندق)"),
        ("/api/hotels/99999", "GET", "طلب فندق برقم غير موجود في النظام", "404 Not Found", "ناجح (إرجاع رسالة غير موجود)"),
        ("/api/hotels", "POST", "إضافة فندق جديد بكافة البيانات الصحيحة", "201 Created", "ناجح (حفظ الفندق بالـ DB)"),
        ("/api/hotels", "POST", "إضافة فندق بدون اسم (حقل إجباري فارغ)", "400 Bad Request", "ناجح (رفض مع رسالة التحقق)"),
        ("/api/rooms", "GET", "تصفية الغرف التابعة لفندق محدد (?hotelId=1)", "200 OK", "ناجح (استرجاع غرف الفندق)"),
        ("/api/bookings", "POST", "إنشاء حجز فندقي بتاريخ مستقبلي صحيح", "201 Created", "ناجح (إنشاء الحجز وحفظه)"),
        ("/api/bookings", "POST", "محاولة حجز مزدوج لغرفة محجوزة بنفس الفترة", "409 Conflict", "ناجح (منع الحجز المزدوج)")
    ]

    for row_idx, data in enumerate(tests_data, start=1):
        row = table_tests.rows[row_idx]
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 70, 70, 90, 90)
            p = cell.paragraphs[0]
            set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            r = p.add_run(val)
            is_result = (col_idx == 4)
            set_run_font(r, font_name=FONT_ARABIC, size_pt=12, bold=is_result, color=COLOR_GREEN if is_result else COLOR_TEXT)

    add_arabic_heading(doc, "2.4 مصفوفة اختبارات التحقق من صحة المدخلات (Input Validation Testing):", level=2)

    table_val = doc.add_table(rows=6, cols=4)
    table_val.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_styling(table_val, border_color="CBD5E1", border_sz="4")

    cols_val = ["الكيان / الحقل البرمجي", "قاعدة التحقق المفروضة", "حالة الفحص السلبي (Negative Test)", "استجابة النظام وتصرفه"]
    for i, c in enumerate(cols_val):
        cell = table_val.rows[0].cells[i]
        set_cell_background(cell, "334155")
        p = cell.paragraphs[0]
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run(c)
        set_run_font(r, font_name=FONT_ARABIC, size_pt=12.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    val_data = [
        ("User.Email", "Required, EmailAddress", "إدخال نص بريد بدون علامة @ أو نطاق", "رفض فوري: يرجى إدخال بريد إلكتروني صالح"),
        ("Hotel.Rating", "Range (1.0 إلى 5.0)", "إدخال تقييم بقيمة 7.5 أو قيمة سالبة", "رفض فوري: التقييم يجب أن يكون بين 1 و 5"),
        ("Room.PricePerNight", "Range (1 إلى 1,000,000)", "إدخال سعر ليلة بقيمة صفر أو سالبة (-100)", "رفض فوري: يجب أن يكون السعر رقماً موجباً أكبر من الصفر"),
        ("Room.Capacity", "Range (1 إلى 20 ضيفاً)", "إدخال سعة استيعابية بقيمة صفر", "رفض فوري: سعة الغرفة يجب أن تتسع لشخص واحد على الأقل"),
        ("Booking.Dates", "قاعدة مخصصة IValidatableObject", "تحديد تاريخ المغادرة مساوياً أو قبل الوصول", "رفض فوري: تاريخ المغادرة يجب أن يكون بعد تاريخ الوصول")
    ]

    for row_idx, data in enumerate(val_data, start=1):
        row = table_val.rows[row_idx]
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 70, 70, 90, 90)
            p = cell.paragraphs[0]
            set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            r = p.add_run(val)
            set_run_font(r, font_name=FONT_ARABIC, size_pt=12, bold=(col_idx == 0), color=COLOR_TEXT)

    add_arabic_heading(doc, "3.4 اختبار منطق الأعمال وحماية العمليات المتزامنة (منع الحجز المزدوج):", level=2)
    add_arabic_paragraph(
        doc,
        "تُعد حماية النظام من الحجوزات المتضاربة (Double Bookings) أهم قاعدة عمل في قطاع الفنادق. "
        "تم فحص هذه القاعدة هندسياً في طبقة الـ Service والـ Repository، حيث يمنع النظام تداخل الفترات الزمنية عبر الاستعلام التالي:\n"
        "• تم تسجيل حجز مسبق للغرفة رقم 101 من تاريخ 10 أكتوبر إلى 15 أكتوبر وتم تأكيده.\n"
        "• تمت محاولة إرسال حجز آخر لنفس الغرفة 101 من تاريخ 12 أكتوبر إلى 18 أكتوبر (فترة متداخلة).\n"
        "• النتيجة الفعلية: رفض النظام الطلب فورياً برمز 409 Conflict مع رسالة خطأ واضحة: 'عذراً، هذه الغرفة محجوزة بالفعل في الفترة المحددة'، "
        "وظلت قاعدة البيانات متسقة بنسبة 100% دون تسجيل أي تضارب."
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # 6. الجزء الرابع: التوليف التراكمي لأسابيع المشروع (W2 - W9)
    # -------------------------------------------------------------
    add_arabic_heading(doc, "5. الجزء الرابع: التوليف التراكمي الشامل لأسابيع المشروع (W2 إلى W9)", level=1)

    add_arabic_paragraph(
        doc,
        "يوضح هذا الجدول التوليفي مطابقة كافة مخرجات المشروع المنفذة مع توزيع الأسابيع والأوزان والدرجات "
        "المعتمدة في دليل مساق التدريب الميداني لجامعة الحكمة:"
    )

    table_weeks = doc.add_table(rows=10, cols=4)
    table_weeks.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_styling(table_weeks, border_color="CBD5E1", border_sz="4")

    cols_w = ["الأسبوع", "المتطلب في الدليل", "الوزن", "ما تم تنفيذه وإنجازه في المشروع بالتفصيل"]
    for i, c in enumerate(cols_w):
        cell = table_weeks.rows[0].cells[i]
        set_cell_background(cell, "1E1B4B")
        p = cell.paragraphs[0]
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        r = p.add_run(c)
        set_run_font(r, font_name=FONT_ARABIC, size_pt=12.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    weeks_summary = [
        ("W2", "مقترح المشروع والمتطلبات (Project Proposal)", "5%", "صياغة وثيقة المقترح، دراسة المشكلة، الأهداف، النطاق، والجدول الزمني، و12 متطلباً وظيفياً و6 متطلبات غير وظيفية."),
        ("W3", "التصميم المعماري ومخططات UML (Clean Architecture)", "15%", "بناء معمارية Clean Architecture ونمط المستودع Repository Pattern، ورسم مخططات Use Case و Sequence و Class Diagrams."),
        ("W4", "تصميم واجهات وتجربة المستخدم (UI/UX Design)", "5%", "هندسة النماذج التفاعلية التوليدية في Adobe XD (XD.xd) بالهوية البصرية الفاخرة (Indigo, Violet, Gold) ودعم الـ RTL."),
        ("W5", "تطوير واجهات البرمجة (Web API Development)", "10%", "بناء 18 نقطة نهاية برمجية مؤمنة بـ JWT مع حماية البيانات ومعالجة الأخطاء وسياسات الـ CORS."),
        ("W6", "فحص واختبار واجهات البرمجة (API Testing)", "5%", "أتمتة الفحص الشامل عبر Postman Collection وتوفير واجهة Swagger UI التفاعلية مع تغطية كافة رموز الاستجابة."),
        ("W7", "تطبيق الويب المتكامل (ASP.NET Core MVC)", "25%", "بناء واجهات MVC بـ Bootstrap 5 RTL، خط Cairo، متغيرات CSS، جداول DataTables.js، ومودالات الحذف التفاعلية."),
        ("W8", "تطبيق الهاتف المحمول (Flutter Mobile App)", "15%", "تطوير تطبيق موبايل متكامل بنمط Provider وقائمة Drawer وشريط BottomNav ونافذة حجز الغرف المنبثقة SheetModal."),
        ("W9", "إدارة الإصدارات والدمج (GitHub Workflow)", "5%", "تطبيق إدارة الفروع والدمج، وفصل تطبيق الموبايل في ريبو مستقل تماماً وفق توجيهات أ.د/ إبراهيم البلطة."),
        ("W10", "التوثيق النهائي والدليل الفني (Documentation)", "15%", "إعداد دليل التثبيت، دليل المستخدم، مصفوفة الاختبارات، الخاتمة والتحديات، وقائمة المراجع العلمية.")
    ]

    for row_idx, data in enumerate(weeks_summary, start=1):
        row = table_weeks.rows[row_idx]
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            r = p.add_run(val)
            is_w = (col_idx == 0)
            is_weight = (col_idx == 2)
            set_run_font(r, font_name=FONT_ARABIC, size_pt=12.5, bold=(is_w or is_weight), color=COLOR_BRAND if is_weight else COLOR_TEXT)

    # -------------------------------------------------------------
    # 7. الجزء الخامس: الخاتمة والتحديات (W10.4 Conclusion)
    # -------------------------------------------------------------
    add_arabic_heading(doc, "6. الجزء الخامس: الخاتمة والتحديات الهندسية والتطويرات المستقبلية (W10.4)", level=1)

    add_arabic_paragraph(
        doc,
        "مثلت تجربة مساق التدريب الميداني الفردي محاكاة واقعية ومتقدمة لبيئات العمل المؤسسية في كبرى شركات هندسة البرمجيات. "
        "تمكن المهندس من قيادة وتطوير كافة مراحل النظام بجهد فردي كامل، والتغلب على العديد من التحديات الهندسية المعقدة:"
    )

    add_card_box(
        doc,
        "التحدي الأول: عزل تطبيق الموبايل في مشروع ومستودع مستقل كلياً (Supervisor Mandate)",
        [
            "• المشكلة: كان تطبيق Flutter مدمجاً كمجلد فرعي ضمن حل الـ .NET، وهو ما يتعارض مع توجيه د/ إبراهيم البلطة بضرورة استقلالية الموبايل.",
            "• الحل الهندسي: تم استخراج مجلد Flutter بالكامل خارج حل الـ Backend، وتأسيس مستودع Git مستقل له بفروعه ودمجه، وتعديل ملفات .gitignore في المشروعين لعزل ملفات البناء الثقيلة، ونشر المشروعين في مستودعين منفصلين على GitHub مع توثيق الربط بينهما."
        ],
        border_color="3730A3",
        bg_color="EEF2FF"
    )

    add_card_box(
        doc,
        "التحدي الثاني: مزامنة التوقيتات القياسية العالمية (UTC Timestamps Synchronization)",
        [
            "• المشكلة: ظهور فوارق زمنية تؤدي لاختلاف يوم الحجز عند إرسال التاريخ من أجهزة الموبايل بسبب التوقيت المحلي للمنطقة الزمنية (Local Timezone Offset).",
            "• الحل الهندسي: توحيد كافة معاملات التواريخ على مستوى طبقة التطبيق (Application Layer) والـ Database بنظام UTC الخالص عبر DateTime.SpecifyKind(d, DateTimeKind.Utc) مما ضمن دقة يوم الحجز والتحقق من التوفر بنسبة 100%."
        ],
        border_color="C9971F",
        bg_color="FFFBEB"
    )

    add_card_box(
        doc,
        "التحدي الثالث: معالجة استجابة شاشات الموبايل وتفادي تجاوز البكسلات (Pixel Overflow Prevention)",
        [
            "• المشكلة: ظهور خطوط التحذير الصفراء والسوداء عند تصغير أو تدوير الشاشة بسبب تجاوز أبعاد المكونات لمساحة الشاشة المتاحة.",
            "• الحل الهندسي: ابتكار عنصر ResponsivePhoneFrame في Flutter مدعوماً بـ LayoutBuilder و SingleChildScrollView مما ضمن تكيف الواجهات تلقائياً مع مختلف قياسات الشاشات دون أي أخطاء تجاوز."
        ],
        border_color="7C3AED",
        bg_color="FAF5FF"
    )

    add_arabic_heading(doc, "خارطة التطويرات المستقبلية للمنظومة (Future Engineering Roadmap):", level=2)
    add_arabic_paragraph(doc, "1. ربط بوابات الدفع الإلكتروني المباشر (Stripe / PayPal / المحافظ الإلكترونية اليمنية مثل ون كاش وجايبي وجوالي).")
    add_arabic_paragraph(doc, "2. تفعيل الإشعارات الفورية السحابية (Push Notifications) عبر Firebase Cloud Messaging لتنبيه النزلاء بمواعيد الدخول والمغادرة.")
    add_arabic_paragraph(doc, "3. دمج محرك توصيات ذكي بالذكاء الاصطناعي لاقتراح الغرف والفنادق المناسبة بناءً على سلوك وتفضيلات النزيل السابقة.")

    # -------------------------------------------------------------
    # 8. الجزء السادس: المراجع العلمية (W10.5 References)
    # -------------------------------------------------------------
    add_arabic_heading(doc, "7. الجزء السادس: المراجع الأكاديمية والتقنية المعتمدة (W10.5 References)", level=1)

    refs_ar = [
        "روبرت سي مارتن (Uncle Bob) — كتاب Clean Architecture: A Craftsman's Guide to Software Structure and Design، دار النشر العالمية Prentice Hall.",
        "شركة مايكروسوفت (Microsoft Corporation) — الوثائق والمراجع الهندسية الرسمية لإطار العمل ASP.NET Core 10.0 و Entity Framework Core، منصة Microsoft Learn.",
        "شركة جوجل (Google LLC) — التوثيق الرسمي ومستندات المطورين لإطار العمل Flutter وإدارة الحالة عبر Provider Pattern، منصة Flutter.dev.",
        "د. روي توماس فيلدينغ (Roy Fielding) — أطروحة المعمارية البرمجية لشبكات الويب ونمط واجهات الـ REST Architectural Style، جامعة كاليفورنيا إيرفاين.",
        "فريق مهندسي الإنترنت الدولي (IETF) — المعيار القياسي الدولي RFC 7519 لرموز التوثيق المشفرة JSON Web Token (JWT).",
        "جامعة الحكمة — دليل مساق التدريب الميداني لمشاريع تطوير البرمجيات الشاملة الفردية (SDLC)، قسم تكنولوجيا المعلومات، صنعاء، الجمهورية اليمنية."
    ]

    for idx, ref in enumerate(refs_ar, start=1):
        p = doc.add_paragraph()
        set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"[{idx}] {ref}")
        set_run_font(r, font_name=FONT_ARABIC, size_pt=13, bold=False, color=COLOR_TEXT)

    # حفظ المستند المكتمل والمصور
    output_path_root = r"E:\Project_ASP.NET\Hotel_Booking_System_W10_Documentation_AR.docx"
    output_path_final = r"E:\Project_ASP.NET\Hotel_Booking_System_W10_Documentation_AR_Final.docx"
    output_path_docs = r"E:\Project_ASP.NET\HOTEL_BOOKIN_SYSTEM\Hotel_Booking_System\docs\Hotel_Booking_System_W10_Documentation_AR.docx"

    # Save to Final path first
    doc.save(output_path_final)
    print("Saved curated report to:", output_path_final)

    try:
        doc.save(output_path_root)
        print("Updated root docx:", output_path_root)
    except PermissionError:
        print("Root docx is locked in Word, saved to Final:", output_path_final)

    try:
        doc.save(output_path_docs)
        print("Updated docs folder docx:", output_path_docs)
    except PermissionError:
        pass

if __name__ == "__main__":
    build_curated_w10_document()
